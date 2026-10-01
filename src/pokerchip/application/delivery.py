"""Human-readable automatic exports plus links to immutable scientific runs."""
from pathlib import Path
import csv,re,shutil,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from ..core.storage import Records,read_json,atomic_json
from .exporting import safe_cell
from .. import __version__

HEADERS=['영상','칩 번호','프레임','실제 시간(s)','중심 x(px)','중심 y(px)','중심 x(m)','중심 y(m)',
         '속도 x(m/s)','속도 y(m/s)','속력(m/s)','가속도 x(m/s²)','가속도 y(m/s²)',
         '회전각(rad)','각속도(rad/s)','각가속도(rad/s²)','속도 x(px/s)','속도 y(px/s)',
         '거리 보정 상태','시간 상태','각도 상태','측정 출처','빈 값 이유','분석 기록']


def trajectory_rows(run):
    # Publish from the already-completed scientific export. Reopening a writable
    # SQLite connection is unnecessary for this read-only presentation step.
    path=run/'export/trajectories.jsonl'
    if path.exists():
        with path.open(encoding='utf-8') as stream:
            for line in stream:
                if line.strip():yield json.loads(line)
    else:
        with Records(run/'records.sqlite',readonly=True) as db:
            yield from db.rows('trajectories')


def human_rows(run,name):
    for r in trajectory_rows(run):
        px=r.get('raw_center_px') or [None,None];xy=r.get('world_center_m') or [None,None]
        yield [name,r['chip_id'],r['frame_index'],r.get('physical_time_s'),*px,*xy,
            r.get('vx_m_s'),r.get('vy_m_s'),r.get('speed_m_s'),r.get('ax_m_s2'),r.get('ay_m_s2'),
            r.get('theta_unwrapped_rad') if r.get('theta_unwrapped_rad') is not None else r.get('theta_wrapped_rad'),
            r.get('omega_rad_s'),r.get('alpha_rad_s2'),r.get('vx_px_s'),r.get('vy_px_s'),
            r.get('geometry_status'),r.get('time_status'),r.get('angle_status'),r.get('source'),r.get('reason'),run.name]

def publish_results(folder,project,issues=()):
    folder=Path(folder);out=folder/'results';out.mkdir(exist_ok=True)
    summary=[]
    with (out/'전체_운동데이터.csv').open('w',encoding='utf-8-sig',newline='') as merged:
        writer=csv.writer(merged);writer.writerow(HEADERS)
        for e in project['experiments']:
            row={'영상':e['name'],'상태':e.get('status'),'칩 수':len(e['participating_chip_ids']),
                 '개수 결정':e.get('count_mode','manual'),'관측 비율':None,'확인할 프레임':None,
                 '거리 보정':None,'실제 시간':None,'초기30프레임 관측':None,'각속도 행':None,'탐색 피팅':None,'격자 긴 방향':e.get('floor_scene',{}).get('long_direction'),
                 '결과 폴더':'','설명':e.get('failure_reason','')}
            # Never present a previous successful run as the result of a failed retry.
            if e.get('last_run') and e.get('status')=='complete':
                run=folder/e['last_run'];q=read_json(run/'export/quality.json')
                leaf=re.sub(r'[^\w가-힣.-]+','_',Path(e['name']).stem)[:60]+'_'+e['id'][-8:]
                dest=out/leaf/run.name;dest.mkdir(parents=True,exist_ok=True)
                for f in (run/'export').iterdir():
                    if f.is_file():shutil.copy2(f,dest/f.name)
                with (dest/'운동데이터.csv').open('w',encoding='utf-8-sig',newline='') as stream:
                    local=csv.writer(stream);local.writerow(HEADERS)
                    for values in human_rows(run,e['name']):
                        values=[safe_cell(v) for v in values];local.writerow(values);writer.writerow(values)
                intake=folder/'intake'/e['id']/'count_preview.jpg'
                if intake.exists():shutil.copy2(intake,dest/'칩개수_확인.jpg')
                with (dest/'확인할_프레임.csv').open('w',encoding='utf-8-sig',newline='') as stream:
                    w=csv.writer(stream);w.writerow(['프레임','칩','확인 이유'])
                    for issue in q.get('review_issues',[]):w.writerow([issue['frame'],issue.get('chip_id',''),issue['reason']])
                row.update({'관측 비율':q.get('observed_fraction_selected_intervals',q['observed_fraction_all_target_frames']),'확인할 프레임':q.get('review_frame_count',0),
                            '초기30프레임 관측':q.get('initial_30_observations'),'각속도 행':q.get('angular_velocity_rows'),
                            '탐색 피팅':'physics_preview.json' if (dest/'physics_preview.json').exists() else '실행 안 됨',
                            '거리 보정':q.get('geometry_status'),'실제 시간':q.get('time_status'),
                            '결과 폴더':str(dest.relative_to(out)),
                            '설명':'자동 검사는 측정 정확도 보증이 아닙니다. 검토 목록과 보정 상태를 확인하세요.'})
                atomic_json(dest/'provenance.json',{'source_run':str(run.relative_to(folder)),
                    'source_hash':e.get('source_hash'),'count_proposal':e.get('count_proposal'),
                    'chip_id_scope':'track number within this video; not physical identity across trials',
                    'rolling_shutter':'not_corrected_by_user_choice'})
            summary.append(row)
    fields=list(summary[0]) if summary else ['영상','상태']
    with (out/'영상별_요약.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        w=csv.DictWriter(stream,fieldnames=fields);w.writeheader()
        for row in summary:w.writerow({k:safe_cell(v) for k,v in row.items()})
    book=Workbook();sheet=book.active;sheet.title='영상별 요약';sheet.append(fields)
    for row in summary:sheet.append([safe_cell(row.get(k)) for k in fields])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='174E61');cell.font=Font(color='FFFFFF',bold=True)
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=28
    book.save(out/'영상별_요약.xlsx')
    atomic_json(out/'summary.json',{'videos':summary,'issues':list(issues),'version':__version__})
    (out/'먼저_읽어주세요.txt').write_text(
        '포커칩 v4.0 분석 결과\n\n전체_운동데이터.csv: 모든 완료 영상의 측정 데이터\n영상별_요약.xlsx: 처리 상태와 결과 위치\n'
        'physics_preview.json과 Farkas_preview 그림: 직접 관측의 탐색 피팅. 공식 확정 계수나 독립 예측 검증이 아닙니다.\n'
        'review_intervals.csv: 원인별로 묶은 확인 구간. 흐림·바닥 경계를 먼저 확인하세요.\n'
        '각 영상 폴더: 운동데이터.csv, trajectories.png, 확인할_프레임.csv, 칩개수_확인.jpg(자동 인식 시)\n\n'
        '실제 시간 = 파일 재생시간 / 8. 기본 촬영 240fps, 저장 30fps입니다.\n'
        'provisional: 보정판 인쇄 치수와 바닥/카메라 고정 가정에 근거한 잠정 거리입니다. 실제 자로 검증하지 않았습니다.\n'
        'declared: 사용자가 지정한 촬영 배속에 따른 시간입니다. 별도 시계 검증은 하지 않았습니다.\n'
        '빈 값은 0이 아닙니다. 가림·충돌 구간·표식 누락·보정 실패 때 속도/각속도를 비울 수 있습니다.\n'
        '각속도는 설정한 최대 회전속도와 프레임 사이 회전 가정에 조건부입니다. 빠른 회전의 alias를 영상만으로 완전히 배제할 수 없습니다.\n'
        '칩 번호는 해당 영상 안의 추적 번호입니다. 다른 영상의 같은 번호가 같은 실물 칩이라는 뜻은 아닙니다.\n'
        'rolling shutter는 보정하지 않습니다. 자동 분석 완료는 실험 정확도 검증 완료가 아닙니다.\n'
        '충돌 후보와 피팅은 별도 검토가 필요합니다. 자동 파이프라인은 불확실한 충돌에 임의 계수를 넣지 않습니다.\n'
        +'\n'.join(issues),encoding='utf-8')
    return out
