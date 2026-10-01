"""Streaming CSV/JSONL, bounded plots, summarized XLSX, second-pass PTS overlay."""
import csv
from fractions import Fraction
from pathlib import Path
import av
import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from ..core.storage import Records,read_json,atomic_json,dumps
from ..measurement.video import frames,metadata

FIELDS={
    "frames":["frame_index","pts","time_base_num","time_base_den","presentation_time_s","physical_time_s","exposure_midpoint_s","exposure_s","source_frame_type","independent_observation","time_status","time_reason","duplicate_suspected","missing_frame_suspected","blur_score","clipped_fraction","rotation_deg_ccw","chip_count_proposal"],
    "observations":["frame_index","chip_id","raw_x_px","raw_y_px","x_m","y_m","radius_px","radius_m","theta_wrapped_rad","theta_sigma_rad","edge_residual_px","visible_arc_fraction","status","identity_status","angle_status","geometry_status","source"],
    "predictions":["frame_index","chip_id","predicted_x_px","predicted_y_px","status","source","last_observed_frame"],
    "manual":["frame_index","chip_id","raw_x_px","raw_y_px","x_m","y_m","radius_px","radius_m","theta_wrapped_rad","theta_sigma_rad","status","identity_status","source"],
    "trajectories":["frame_index","chip_id","physical_time_s","raw_x_px","raw_y_px","x_m","y_m","theta_wrapped_rad","theta_unwrapped_rad","vx_m_s","vy_m_s","ax_m_s2","ay_m_s2","omega_rad_s","alpha_rad_s2","direction_rad","speed_m_s","vx_px_s","vy_px_s","ax_px_s2","ay_px_s2","geometry_status","time_status","status","angle_status","source","segment","fit_samples","window_s","actual_window_s","reason"],
    "events":["id","kind","status","fit_eligible","frame_start","frame_end","closest_frame","time_s","time_sigma_s","a_m_s","c_m_s","c_after_m_s","e_n_obs","e_t_obs","impact_parameter_m","incidence_rad","reason"]}

for _table in ('manual','trajectories'):
    FIELDS[_table]+= ['observable','fit_enabled','scope_reason','scope_segment','position_status','orientation_status']
FIELDS['events']+=['candidate_frame_interval','contact_frame_interval','boundary_frame_interval','unusable_frame_intervals','review_basis','review_reason']


def flatten(row):
    result=dict(row)
    for key,names in [("raw_center_px",("raw_x_px","raw_y_px")),("world_center_m",("x_m","y_m")),("predicted_center_px",("predicted_x_px","predicted_y_px"))]:
        value=row.get(key)
        for i,name in enumerate(names):result[name]=value[i] if value is not None else None
    return result


def safe_cell(value):
    # User names/reasons are data, never spreadsheet formulas.
    if isinstance(value,(dict,list,tuple)):return dumps(value)
    if isinstance(value,str) and value.startswith(("=","+","-","@")):return "'"+value
    return value


def table_csv(path,fields,rows):
    with Path(path).open("w",encoding="utf-8-sig",newline="") as stream:
        writer=csv.DictWriter(stream,fieldnames=fields,extrasaction="ignore")
        writer.writeheader()
        for row in rows:writer.writerow({k:safe_cell(v) for k,v in flatten(row).items()})


def export_run(run):
    run=Path(run);config=read_json(run/"effective_settings.json");manifest=read_json(run/"manifest.json")
    experiment=read_json(run/"experiment.json")
    out=run/"export";out.mkdir(exist_ok=True)
    with Records(run/"records.sqlite",readonly=True) as db:
        for table,fields in FIELDS.items():
            table_csv(out/(table+".csv"),fields,db.rows(table))
            with (out/(table+".jsonl")).open("w",encoding="utf-8") as stream:
                for row in db.rows(table):stream.write(dumps(row)+"\n")
        frame_count=db.count("frames");observed=db.count("observations")
        eligible=0;kinematic=0;angles=0;manual=0;low=0
        for row in db.rows("trajectories"):
            kinematic+=row.get("vx_m_s") is not None;angles+=row.get("omega_rad_s") is not None;manual+=row.get("source")=="manual_corrected"
        for row in db.rows("observations"):low+=row["status"]=="low_confidence"
        events=list(db.rows("events"));eligible=sum(e.get("fit_eligible",False) for e in events)
        denominator=frame_count*len(experiment["participating_chip_ids"])
        quality={"frames":frame_count,"expected_chip_frames":denominator,"observation_rows":observed,
                 "observed_fraction_all_target_frames":observed/denominator if denominator else 0,
                 "physical_velocity_rows":kinematic,"angular_velocity_rows":angles,"manual_rows":manual,
                 "low_confidence_observations":low,"event_candidates":len(events),"fit_eligible_events":eligible,
                 "accuracy_status":"not_independently_validated", "missing_evidence":[]}
        if config["mode"]!="synthetic_demo":quality["missing_evidence"].append("실제 240fps·새 두 표식 동적 정확도 검증 대기")
        if config["time_profile"]["status"] not in ("verified","synthetic_known_clock"):quality["missing_evidence"].append("물리 시간축 독립 검증 대기")
        if config["calibration"].get("status") not in ("verified","synthetic"):quality["missing_evidence"].append("물리 좌표 보정 검증 대기")
        if any(c.get("mass_kg") is None or (c.get("inertia_kg_m2") is None and c.get("inertia_model")!="uniform_disk") for c in config["chips"] if c["id"] in experiment["participating_chip_ids"]):quality["missing_evidence"].append("질량/관성 실측 또는 명시적 근사 선택 대기")
        from ..analysis.quality import enrich_quality
        quality=enrich_quality(db,quality,config,experiment)
        quality['measurement_status']='review_required' if quality['review_frame_count'] or not angles else 'provisional_measurements_available'
        quality['initial_30_observations']=sum(1 for r in db.rows('observations') if r['frame_index']<30)
        quality['blur_review_rows']=sum(bool(r.get('measurement_warning')) for r in db.rows('observations'))
        quality['floor_boundary_rows']=sum(r.get('surface_status')=='floor_boundary' for r in db.rows('observations'))
        grouped={}
        for issue in quality.get('review_issues',[]):
            key=(issue.get('chip_id',''),issue['reason']);grouped.setdefault(key,[]).append(issue['frame'])
        spans=[]
        for (chip,reason),frames_ in grouped.items():
            for f in sorted(set(frames_)):
                if spans and spans[-1]['chip_id']==chip and spans[-1]['reason']==reason and f==spans[-1]['end_frame']+1:spans[-1]['end_frame']=f
                else:spans.append({'chip_id':chip,'reason':reason,'start_frame':f,'end_frame':f})
        table_csv(out/'review_intervals.csv',['chip_id','reason','start_frame','end_frame'],spans)
        atomic_json(out/"quality.json",quality)
        atomic_json(out/"events.json",events)
        atomic_json(out/"parameters.json",{"source":"configured_not_fitted","parameters":config["physics"],"fit_status":"not_run_use_fit_dataset"})
        table_csv(out/"quality.csv",["metric","value"],({"metric":k,"value":v} for k,v in quality.items() if not isinstance(v,(list,dict))))
        table_csv(out/"parameters.csv",["parameter","value","status"],({"parameter":k,"value":v,"status":"configured_not_fitted"} for k,v in config["physics"].items()))
        atomic_json(out/"corrections.json",{"history":config["corrections"],"cursor":config["correction_cursor"],"audit":config.get("correction_audit",[])})
        workbook=Workbook();workbook.remove(workbook.active)
        sheets={"experiments":(["experiment_id","name","mode","frames","raw_data"],[[experiment["id"],experiment["name"],config["mode"],frame_count,"trajectories.csv"]]),
                "chips":(["id","radius_m","radius_status","mass_kg","thickness_m","inertia_kg_m2","inertia_model"],[[c.get(k) for k in ("id","radius_m","radius_status","mass_kg","thickness_m","inertia_kg_m2","inertia_model")] for c in config["chips"]]),
                "events":(FIELDS["events"],[[e.get(k) for k in FIELDS["events"]] for e in events]),
                "parameters":(["parameter","value","status"],[[k,v,"configured_not_fitted"] for k,v in config["physics"].items()]),
                "quality":(["metric","value"],[[k,v] for k,v in quality.items() if not isinstance(v,(list,dict))]),
                "data_dictionary":(["table","column","unit_or_definition"],[[table,k,unit(k)] for table,fields in FIELDS.items() for k in fields])}
        for name,(header,rows) in sheets.items():
            sheet=workbook.create_sheet(name);sheet.append(header)
            for row in rows:sheet.append([safe_cell(v) for v in row])
            sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="214B65")
            for column in sheet.columns:
                letter=column[0].column_letter;sheet.column_dimensions[letter].width=min(50,max(16,len(str(column[0].value))+3))
        workbook.save(out/"summary.xlsx")
        plot_trajectories(db,out/"trajectories.png")
    limitations="\n".join("- "+x for x in quality["missing_evidence"]) or "- 합성 fixture에만 유효합니다. 실제 실험 검증을 뜻하지 않습니다."
    (out/"REPORT_KO.md").write_text(f"# 포커칩 분석 보고서\n\n모드: **{config['mode']}**\n\n실험: {experiment['name']} · run: {manifest['id']}\n\n"
        f"전체 {frame_count}프레임, 참여 칩 기준 {denominator}개 관측 기회 중 {observed}개 관측 행입니다. 검출률은 독립 정답 recall이 아닙니다.\n\n"
        f"물리 속도 {kinematic}행, 각속도 {angles}행, 사건 후보 {len(events)}개, 승인되어 피팅 가능한 사건 {eligible}개.\n\n"
        "원시 관측은 observations, 추적 예측은 predictions, 수동 반영 관측은 manual, 국소 미분은 trajectories에 구분했습니다. 빈 수치는 0이 아닙니다.\n\n"
        f"시간 설정: {config['time_profile']['status']} · 구간별 slow_factor로 PTS를 실제 시간으로 변환합니다. declared는 촬영조건에 근거한 잠정값입니다.\n\n"
        f"검토 프레임 {quality.get('review_frame_count',0)}개 · ID 모호 관측 {quality.get('ambiguous_id_rows',0)}개.\n\n"
        "## 제한 및 미확정 근거\n\n"+limitations+"\n\n모든 수치는 원본·설정·수정 이력·코드 hash에 연결됩니다. profile을 수정한 뒤 새 run을 생성하세요.\n",encoding="utf-8")
    if config.get('quick_setup',{}).get('preview_physics',False) and experiment.get('quick_workflow',False):
        from ..models.preview_physics import preview
        preview(run,config,experiment)
    return out


def unit(key):
    if key.endswith("_m_s2"):return "m/s²"
    if key.endswith("_m_s"):return "m/s"
    if key.endswith("_rad_s2"):return "rad/s²"
    if key.endswith("_rad_s"):return "rad/s, CCW+"
    if key.endswith("_rad"):return "rad, world CCW+"
    if key.endswith("_px"):return "px, image y down"
    if key.endswith("_m"):return "m, world y up"
    if key.endswith("_s"):return "s; physical and presentation separate"
    return "상태/식별자/무차원; docs/DATA_DICTIONARY_KO.md 참조"


def plot_trajectories(db,path):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    stride=max(1,db.count("trajectories")//4000)
    rows=[r for i,r in enumerate(db.rows("trajectories")) if i%stride==0]
    fig,axes=plt.subplots(2,2,figsize=(11,7),layout="constrained")
    for chip in sorted({r["chip_id"] for r in rows}):
        selected=[r for r in rows if r["chip_id"]==chip]
        metric=all(r.get("world_center_m") is not None for r in selected)
        xy=np.array([r["world_center_m"] if metric else r["raw_center_px"] for r in selected])
        axes[0,0].plot(xy[:,0],xy[:,1],".",ms=2,label=chip)
        frame=[r["frame_index"] for r in selected]
        for ax,key in [(axes[0,1],"theta_wrapped_rad"),(axes[1,0],"vx_m_s"),(axes[1,1],"omega_rad_s")]:
            ax.plot(frame,[r.get(key) if r.get(key) is not None else np.nan for r in selected],".",ms=2,label=chip)
    for ax,title in zip(axes.flat,["Observed position (m if calibrated, otherwise px)","Wrapped angle (rad)","vx (m/s) | blank = gated","omega (rad/s) | blank = gated"]):
        ax.set_title(title,fontsize=10);ax.grid(alpha=.25)
        if rows:ax.legend(fontsize=7)
    for ax in (axes[0,1],axes[1,0],axes[1,1]):ax.set_xlabel("decoded frame index")
    fig.savefig(path,dpi=140);plt.close(fig)


def overlay(run,source,destination=None):
    run=Path(run);destination=Path(destination or run/"export/overlay.mp4")
    partial=destination.with_name(destination.stem+".partial.mp4")
    exp=read_json(run/"experiment.json")
    meta=metadata(source)
    rate=Fraction(meta["average_rate"]) if meta["average_rate"]!="None" else Fraction(30)
    with Records(run/"records.sqlite",readonly=True) as db, av.open(str(partial),"w") as container:
        stream=None
        with destination.with_suffix(".time_map.csv").open("w",encoding="utf-8-sig",newline="") as text:
            writer=csv.writer(text);writer.writerow(["frame_index","source_pts","source_time_base","presentation_s","physical_s","overlay_pts"])
            origin=None
            for timing,image in frames(source,start=exp["interval"][0],end=exp["interval"][1]):
                f=timing["frame_index"]
                records=list(db.rows("manual",start=f,end=f));ft=next(db.rows("frames",start=f,end=f),timing)
                for row in records:
                    center=tuple(np.round(row["raw_center_px"]).astype(int));radius=int(row["radius_px"])
                    color=(50,200,40) if row.get("source")!="manual_corrected" else (0,200,255)
                    cv2.circle(image,center,radius,color,2);cv2.putText(image,row["chip_id"],(center[0]-radius,center[1]-radius-5),cv2.FONT_HERSHEY_SIMPLEX,.6,color,2)
                label=f"frame {f} | presentation {timing['presentation_time_s']} s | physical {ft.get('physical_time_s')} s"
                cv2.putText(image,label,(16,28),cv2.FONT_HERSHEY_SIMPLEX,.55,(30,30,255),2)
                if stream is None:
                    stream=container.add_stream("libx264",rate=rate);stream.width=image.shape[1];stream.height=image.shape[0];stream.pix_fmt="yuv420p"
                    stream.options={"crf":"20","preset":"fast"}
                if timing["pts"] is None:raise ValueError("PTS 없는 overlay는 정확한 재생시간을 보존할 수 없습니다.")
                tb=Fraction(timing["time_base_num"],timing["time_base_den"])
                if origin is None:origin=timing["pts"]
                frame=av.VideoFrame.from_ndarray(image,format="bgr24");frame.pts=timing["pts"]-origin;frame.time_base=tb
                for packet in stream.encode(frame):container.mux(packet)
                writer.writerow([f,timing["pts"],str(tb),timing["presentation_time_s"],ft.get("physical_time_s"),frame.pts])
            if stream:
                for packet in stream.encode():container.mux(packet)
    partial.replace(destination)
    return destination


def summary_batch(runs,destination):
    workbook=Workbook();sheet=workbook.active;sheet.title="experiments"
    header=["run_id","experiment_id","mode","frames","observation_rows","physical_velocity_rows","event_candidates","report"]
    sheet.append(header)
    for run in runs:
        run=Path(run);m=read_json(run/"manifest.json");q=read_json(run/"export/quality.json")
        sheet.append([m["id"],m["experiment_id"],m["mode"],q["frames"],q["observation_rows"],q["physical_velocity_rows"],q["event_candidates"],str(run/"export/REPORT_KO.md")])
    sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
    workbook.save(destination)
