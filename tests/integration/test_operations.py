import copy
import json
from pathlib import Path
import threading
import time
import numpy as np
import pytest
from openpyxl import load_workbook
from pokerchip.config import create_project,default_project,register,save_project,relink
from pokerchip.storage import Records,read_json,atomic_json,correction,undo_redo,active_corrections
from pokerchip.demo import make_video,make_demo,fit_fixture
from pokerchip.pipeline import analyze,Cancelled,stage_keys
from pokerchip.jobs import Batch,Control
from pokerchip.video import frames
from pokerchip.fitting import fit_free,fit_impacts,bootstrap
from pokerchip.selection import confirm_start
from pokerchip.validation import split_trials


def test_pending_end_to_end_exports_and_cache(tmp_path):
    folder=tmp_path/"한글 공백";p=create_project(folder);p["analysis"]["radius_px"]=[28,36]
    p["time_profile"]["status"]="pending"
    video=make_video(tmp_path/"원본 영상.mp4",count=8,kind="single")
    e=register(folder,p,[video],["chip_1"])[0]
    confirm_start(e,video,0)
    result=analyze(folder,p,e);run=Path(result["run"])
    with Records(run/"records.sqlite") as db:
        assert db.count("frames")==8
        assert db.count("observations")>=6
        assert all(r["physical_time_s"] is None and r["omega_rad_s"] is None and r["vx_m_s"] is None for r in db.rows("trajectories"))
    assert analyze(folder,p,e)["cached"]
    workbook=load_workbook(run/"export/summary.xlsx",read_only=True)
    assert set(workbook.sheetnames)=={"experiments","chips","events","parameters","quality","data_dictionary"}
    assert (run/"export/trajectories.png").stat().st_size>1000
    assert read_json(run/"manifest.json")["status"]=="complete"
    old=stage_keys(p,e,"hash");p["time_profile"].update(status="verified",evidence="test clock")
    new=stage_keys(p,e,"hash");assert old["raw"]==new["raw"] and old["analysis"]!=new["analysis"]
    p["chips"][0]["mass_kg"]=.01
    newer=stage_keys(p,e,"hash");assert new["raw"]==newer["raw"] and new["analysis"]==newer["analysis"] and new["physics"]!=newer["physics"]
    p["plot"]["color"]="red";styled=stage_keys(p,e,"hash");assert styled["analysis"]==newer["analysis"] and styled["export"]!=newer["export"]


def test_cancel_resume_equals_clean_and_fail_isolation(tmp_path):
    folder=tmp_path/"project";p=create_project(folder);p["analysis"].update(radius_px=[28,36],chunk_frames=2)
    video=make_video(tmp_path/"a.mp4",count=10,kind="single");e=register(folder,p,[video],["chip_1"])[0]
    confirm_start(e,video,0)
    control=Control()
    def callback(message):
        if message["stage"]=="detect":control.cancelled.set()
    with pytest.raises(Cancelled):analyze(folder,p,e,control,callback)
    checkpoints=list((folder/"cache").glob("*/checkpoint.json"));assert checkpoints
    assert read_json(checkpoints[0])["last_frame"]==1
    resumed=analyze(folder,p,e)
    cleanfolder=tmp_path/"clean";cleanfolder.mkdir();clean=analyze(cleanfolder,p,e)
    with Records(Path(resumed["run"])/"records.sqlite") as a,Records(Path(clean["run"])/"records.sqlite") as b:
        assert list(a.rows("observations"))==list(b.rows("observations"))
        assert a.count("frames")==b.count("frames")==10
    broken=tmp_path/"broken.mp4";broken.write_bytes(b"invalid codec container")
    added=register(folder,p,[broken,tmp_path/"missing.mp4",video],["chip_1"])
    confirm_start(added[-1],video,0)
    results=Batch(folder,p).run()
    assert [r["status"] for r in results]==["complete","failed","failed","complete"]


def test_pause_resumes_and_correction_audit(tmp_path):
    control=Control();control.paused.set();finished=threading.Event()
    t=threading.Thread(target=lambda:(control.check(),finished.set()));t.start()
    assert not finished.wait(.12);control.paused.clear();assert finished.wait(1);t.join()
    p=default_project()
    correction(p,"center",{"frames":[1,1],"chip_id":"chip_1"},{"point_px":[1,2]},"visible boundary")
    undo_redo(p,-1);assert active_corrections(p)==[]
    undo_redo(p,1);assert len(active_corrections(p))==1
    undo_redo(p,-1);correction(p,"angle",{},0,"new reason")
    assert len(p["abandoned_corrections"])==1 and len(p["correction_audit"])==5


def test_atomic_write_failure_never_replaces_complete(tmp_path,monkeypatch):
    import pokerchip.storage as storage
    path=tmp_path/"state.json";atomic_json(path,{"status":"old"})
    def diskfull(*args):raise OSError("simulated disk full / permission denied")
    monkeypatch.setattr(storage.os,"replace",diskfull)
    with pytest.raises(OSError):atomic_json(path,{"status":"complete"})
    assert read_json(path)=={"status":"old"}
    assert not list(tmp_path.glob("*.partial-*"))


def test_equal_basename_and_hash_relink(tmp_path):
    folder=tmp_path/"project";p=create_project(folder)
    a=tmp_path/"one"/"동일 이름.mp4";b=tmp_path/"two"/"동일 이름.mp4"
    a.parent.mkdir();b.parent.mkdir();a.write_bytes(b"a");b.write_bytes(b"b")
    ea,eb=register(folder,p,[a,b]);assert ea["id"]!=eb["id"]
    from pokerchip.storage import file_hash
    ea["source_hash"]=file_hash(a)
    with pytest.raises(ValueError):relink(folder,p,ea["id"],b)
    moved=tmp_path/"moved.mp4";moved.write_bytes(b"a");relink(folder,p,ea["id"],moved)
    assert ea["video_uri"]==str(moved.resolve())


def test_independent_parameter_recovery_and_holdout():
    data=fit_fixture();free=data["free_trials"][:2]
    fit=fit_free(free,starts=(.1,),max_nfev=60)
    assert abs(fit["parameters"]["mu_bottom"]-.12)<.003
    collision=fit_impacts(data["impact_trials"][:12],eiv=False,max_nfev=100)
    for key,tolerance in [("e_normal",.015),("e_tangential",.025),("mu_collision",.015)]:assert abs(collision["parameters"][key]-data["truth"][key])<tolerance
    assert collision["diagnostics"]["identifiability"]=="identified_locally"
    fixed=fit_impacts(data["impact_trials"][:12],eiv=False,max_nfev=100,fixed_e_normal=.61)
    assert fixed["parameters"]["e_normal"]==.61
    assert fixed["normal_scope"]=="fixed_from_separate_position_based_normal_stage"
    narrow=fit_impacts([copy.deepcopy(data["impact_trials"][0]) for _ in range(5)],eiv=False,max_nfev=60)
    assert narrow["diagnostics"]["identifiability"]!="identified_locally"
    train,held=split_trials(data["free_trials"],["session_0","session_1"],["session_2"])
    assert not {t["id"] for t in train}&{t["id"] for t in held}
    with pytest.raises(ValueError):split_trials(data["free_trials"],["session_0"],["session_0"])
    data["free_trials"][0]["time_status"]="pending"
    with pytest.raises(ValueError):fit_free(data["free_trials"])
