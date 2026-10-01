import copy
import numpy as np
import pytest
from pokerchip.core.config import default_project
from pokerchip.core.review import (scope,set_endpoint,validate_boundary,event_review_basis,
    event_identity,event_source_identity,validate_intervals)
from pokerchip.analysis.kinematics import event_barriers,kinematic_at,apply_corrections,suggest_boundary
from pokerchip.models.study import free_trials
from pokerchip.core.task import task_scope,least_squares
from pokerchip.application.jobs import Control
from pokerchip.application.pipeline import Cancelled


def rows(n=90):
    return [dict(chip_id='chip_1',frame_index=f,physical_time_s=f/240,status='observed',raw_center_px=[f,0],
        world_center_m=[f/240,0],radius_px=20,theta_wrapped_rad=f/240,theta_sigma_rad=.01,
        omega_rad_s=1.,speed_m_s=1.,covariance_world=(np.eye(3)*1e-8).tolist()) for f in range(n)]


def test_actual_collision_excluded_from_fit_still_splits_velocity():
    cfg=default_project()['analysis'];data=rows()
    for r in data:
        if r['frame_index']>=32:r['world_center_m'][0]=31.5/240-(r['frame_index']-31.5)/240
    ev=dict(pair=['chip_1','chip_2'],kind='isolated_binary',status='excluded',frame_start=28,frame_end=36,contact_frame_interval=[31,32])
    barriers=event_barriers([ev],'chip_1')
    pre=kinematic_at(data[31],data,cfg,barriers);post=kinematic_at(data[32],data,cfg,barriers)
    assert pre['vx_m_s']==pytest.approx(1) and post['vx_m_s']==pytest.approx(-1)
    assert pre['fit_frame_interval'][1]==31 and post['fit_frame_interval'][0]==32


def test_free_fit_never_rejoins_adjacent_frames_across_collision():
    cfg=default_project();cfg['analysis']['omega_bound_rad_s']=500.;e={'id':'a','participating_chip_ids':['chip_1'],'session_id':'s'}
    ev=dict(pair=['chip_1','chip_2'],kind='isolated_binary',status='excluded',frame_start=40,frame_end=45,contact_frame_interval=[44,45])
    trials=free_trials(cfg,e,rows(100),[ev],True)
    assert len(trials)==2
    assert all(not (min(t['initialization_frames'])<=44 and max(t['frames'])>=45) for t in trials)


def test_stop_and_exit_have_different_observation_semantics():
    e={'participating_chip_ids':['chip_1','chip_2'],'interval':[3,100]}
    set_endpoint(e,'chip_1',40,'stopped');assert scope(e,'chip_1',60)['observable']
    set_endpoint(e,'chip_1',50,'exit');assert scope(e,'chip_1',50)['observable']
    assert not scope(e,'chip_1',51)['observable'] and scope(e,'chip_2',51)['observable']
    set_endpoint(e,'chip_1',70,'reentry');assert scope(e,'chip_1',70)['observable']
    assert not scope(e,'chip_1',69)['observable']


def test_surface_exit_preserves_positions_but_disables_friction():
    e={'participating_chip_ids':['chip_1'],'interval':[0,None]};set_endpoint(e,'chip_1',30,'surface_exit')
    assert scope(e,'chip_1',31)['observable'] and not scope(e,'chip_1',31)['fit_enabled']


def test_repeated_departures_preserve_gaps_and_failed_edits_are_atomic():
    e={'participating_chip_ids':['chip_1'],'interval':[0,100]}
    set_endpoint(e,'chip_1',30,'exit');set_endpoint(e,'chip_1',50,'reentry');set_endpoint(e,'chip_1',70,'exit')
    assert not scope(e,'chip_1',40)['observable'] and scope(e,'chip_1',60)['observable']
    assert not scope(e,'chip_1',71)['observable']
    before=copy.deepcopy(e)
    with pytest.raises(ValueError):set_endpoint(e,'chip_1',40,'exit')
    assert e==before


def test_surface_departure_after_reentry_preserves_earlier_invisible_gap():
    e={'participating_chip_ids':['chip_1'],'interval':[0,100]}
    set_endpoint(e,'chip_1',20,'exit');set_endpoint(e,'chip_1',40,'reentry');set_endpoint(e,'chip_1',60,'surface_exit')
    assert not scope(e,'chip_1',30)['observable']
    assert scope(e,'chip_1',60)['fit_enabled']
    assert scope(e,'chip_1',61)['observable'] and not scope(e,'chip_1',61)['fit_enabled']


def test_boundary_rejects_unusable_reference_frames():
    assert validate_boundary(31,32)==[31,32]
    with pytest.raises(ValueError):validate_boundary(31,32,[[30,31]])


def test_center_recomputes_angle_from_real_rim_coordinate():
    cfg=default_project();r=rows(1)[0];r['raw_center_px']=[0,0]
    r['markers']={'chip_1':{'rim':{'point_px':[20,0]},'rim_ambiguous':False}}
    edit={'id':'m','action':'center','target':{'chip_id':'chip_1','frames':[0,0]},'after':{'point_px':[0,5]}}
    cal={'status':'verified','H':[[.001,0,0],[0,-.001,0],[0,0,1]]}
    q=apply_corrections(r,[edit],cal,cfg['chips'])
    assert q['theta_wrapped_rad']==pytest.approx(np.arctan2(5,20))
    assert q['orientation_status']=='center_recomputed'


def test_cancel_reaches_optimizer_evaluations():
    control=Control();calls=[]
    def residual(x):calls.append(1);control.cancel();return x-1
    with task_scope(control=control):
        with pytest.raises(Cancelled):least_squares(residual,[0.])
    assert len(calls)==1


def test_change_point_refines_wide_candidate_to_adjacent_frames():
    data=rows(70);other=copy.deepcopy(data)
    for r in data:r['world_center_m'][0]=(r['frame_index']/240 if r['frame_index']<=31 else 31/240)
    for r in other:
        r['chip_id']='chip_2';r['world_center_m'][0]=31/240+.04+(0 if r['frame_index']<=31 else (r['frame_index']-31)/240)
    ev={'pair':['chip_1','chip_2'],'frame_start':28,'frame_end':36}
    def get(chip,a,b):return [r for r in data+other if r['chip_id']==chip and a<=r['frame_index']<=b]
    proposal=suggest_boundary(ev,get)
    assert proposal in ([30,31],[31,32])  # exact kink sample belongs to either continuous fit


def test_review_invalidates_after_measurement_change():
    cfg=default_project();ev={'pair':['chip_1','chip_2'],'frame_start':20,'frame_end':24}
    rr=rows();before=event_review_basis(ev,rr,cfg);rr[21]['world_center_m'][0]+=.001
    assert event_review_basis(ev,rr,cfg)!=before


def test_first_run_and_reanalysis_use_same_event_identity():
    event={'pair':['chip_2','chip_1'],'closest_frame':42}
    first={'id':'exp_1','source_hash':None,'start_selection':{'source_hash':'video_sha256'}}
    later={**first,'source_hash':'video_sha256'}
    assert event_identity(event_source_identity(first),event)==event_identity(event_source_identity(later),event)


def test_chip_intervals_stay_inside_analysis_interval():
    experiment={'participating_chip_ids':['chip_1'],'interval':[10,20],
        'chip_intervals':{'chip_1':[{'start':10,'end':20,'fit_enabled':True}]}}
    assert validate_intervals(experiment) is experiment
    experiment['chip_intervals']['chip_1'][0]['end']=21
    with pytest.raises(ValueError):validate_intervals(experiment)


def test_excel_accepts_nested_collision_boundary(tmp_path):
    from openpyxl import Workbook,load_workbook
    from pokerchip.application.exporting import safe_cell
    import json
    book=Workbook();book.active.append([safe_cell([31,32]),safe_cell({'pre':[1,2]})]);path=tmp_path/'events.xlsx';book.save(path)
    cells=next(load_workbook(path,read_only=True).active.values)
    assert json.loads(cells[0])==[31,32] and json.loads(cells[1])=={'pre':[1,2]}


def test_selected_interval_coverage_does_not_count_frames_after_exit():
    from pokerchip.analysis.quality import enrich_quality
    cfg=default_project();e={'participating_chip_ids':['chip_1'],'interval':[0,9]}
    set_endpoint(e,'chip_1',4,'exit')
    class DB:
        def rows(self,table,start=None,end=None):
            if table=='frames':return iter([{'frame_index':i,'presentation_time_s':i/30} for i in range(10)])
            return iter([dict(frame_index=start,chip_id='chip_1',raw_center_px=[0,0],status='observed' if start<=4 else 'outside_interval')])
    q=enrich_quality(DB(),{'observed_fraction_all_target_frames':.5},cfg,e)
    assert q['observed_fraction_all_target_frames']==.5
    assert q['observed_fraction_selected_intervals']==1 and q['expected_visible_chip_frames']==5


def test_bootstrap_normal_velocity_uncertainty_uses_velocity_units():
    from pokerchip.models.fitting import bootstrap
    trials=[dict(id=str(i),session_id=str(i),pre=[[0,0,2,0,0,0],[0,0,0,0,0,0]],
                 post=[[0,0,0,0,0,0],[0,0,2,0,0,0]],normal_sigma=.1,
                 shared_scale_sigma_fraction=.2,shared_clock_sigma_fraction=.1) for i in range(3)]
    checked=[]
    def fit(sample):
        for tr in sample:assert tr['normal_sigma']/tr['pre'][0][2]==pytest.approx(.05)
        checked.append(1)
        return {'diagnostics':{'optimizer_success':True,'identifiability':'identified_locally','active_bounds':[]},
                'physical_status':'admissible','parameters':{'e_normal':.8}}
    result=bootstrap(trials,fit,count=3)
    assert len(checked)==3 and result['count']==3
